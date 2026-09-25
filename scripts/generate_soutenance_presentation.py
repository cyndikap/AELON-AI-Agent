# -*- coding: utf-8 -*-
"""Generate the AELON 'soutenance' storytelling deck (18 slides).

Narrative: Probleme metier -> Donnee -> IA -> Agents -> Decision.
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BLUE = RGBColor(0x0A, 0x3D, 0x91)
GOLD = RGBColor(0xD4, 0xAF, 0x37)
GREY = RGBColor(0x66, 0x70, 0x85)
BG = RGBColor(0xF5, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1B, 0x22, 0x33)
GREEN = RGBColor(0x1E, 0x7A, 0x4C)
AMBER = RGBColor(0xB8, 0x7A, 0x00)
LIGHT_LINE = RGBColor(0xDD, 0xE3, 0xEC)
ZEBRA = RGBColor(0xEE, 0xF1, 0xF6)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL_SLIDES = 18

OUTPUT_PATH = Path(__file__).resolve().parent.parent / "docs" / "AELON_Soutenance_Presentation.pptx"


# --------------------------------------------------------------------------- #
# Low-level helpers
# --------------------------------------------------------------------------- #
def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    tree = slide.shapes._spTree
    tree.remove(bg._element)
    tree.insert(2, bg._element)
    return slide


def add_textbox(slide, left, top, width, height, text, size=16, color=DARK,
                 bold=False, italic=False, align=PP_ALIGN.LEFT, font="Segoe UI"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return box


def add_header(slide, kicker, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    bar.shadow.inherit = False
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(0.38), Inches(0.28), Inches(0.28))
    dot.fill.solid()
    dot.fill.fore_color.rgb = GOLD
    dot.line.fill.background()
    dot.shadow.inherit = False
    add_textbox(slide, Inches(0.92), Inches(0.06), Inches(11.5), Inches(0.32), kicker,
                size=12, color=GOLD, bold=True)
    add_textbox(slide, Inches(0.92), Inches(0.37), Inches(11.5), Inches(0.6), title,
                size=24, color=WHITE, bold=True)


def add_footer(slide, page_no):
    add_textbox(slide, Inches(0.5), Inches(7.14), Inches(8), Inches(0.3),
                "AELON — Donnee -> Connaissance -> IA -> Decision", size=9.5, color=GREY, italic=True)
    add_textbox(slide, Inches(12.1), Inches(7.14), Inches(0.8), Inches(0.3),
                f"{page_no}/{TOTAL_SLIDES}", size=9.5, color=GREY, align=PP_ALIGN.RIGHT)


def add_bullets(slide, left, top, width, height, items, size=15, color=DARK, header_color=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        level = 0
        text = item
        if isinstance(item, tuple):
            text, level = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        p.line_spacing = 1.2
        p.space_after = Pt(8)
        run = p.add_run()
        prefix = "• " if level == 0 else "— "
        run.text = f"{prefix}{text}"
        run.font.size = Pt(size if level == 0 else size - 1)
        run.font.color.rgb = (header_color or color) if level == 0 else GREY
        run.font.name = "Segoe UI"
    return box


def add_quote(slide, left, top, width, height, text, size=17):
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.06), height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = GOLD
    accent.line.fill.background()
    accent.shadow.inherit = False
    add_textbox(slide, left + Inches(0.25), top, width - Inches(0.25), height,
                text, size=size, color=BLUE, italic=True, bold=True)


def add_placeholder(slide, left, top, width, height, label):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xE6, 0xEA, 0xF2)
    box.line.color.rgb = LIGHT_LINE
    box.line.dash_style = None
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size = Pt(13)
    run.font.italic = True
    run.font.color.rgb = GREY


# --------------------------------------------------------------------------- #
# Flow diagram helpers
# --------------------------------------------------------------------------- #
def add_flow_horizontal(slide, steps, top, left=Inches(0.5), total_width=Inches(12.3),
                         box_h=Inches(0.85), font_size=10.5, gap=Inches(0.15)):
    n = len(steps)
    box_w = (total_width - gap * (n - 1)) / n
    for i, step in enumerate(steps):
        x = left + i * (box_w + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, box_w, box_h)
        box.fill.solid()
        box.fill.fore_color.rgb = BLUE if i % 2 == 0 else GOLD
        box.line.fill.background()
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = step
        run.font.size = Pt(font_size)
        run.font.bold = True
        run.font.color.rgb = WHITE if i % 2 == 0 else DARK
        if i < n - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + box_w, top + box_h / 2 - Inches(0.15),
                                            gap, Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GREY
            arrow.line.fill.background()
            arrow.shadow.inherit = False
    return top + box_h


def add_flow_vertical(slide, steps, center_x, top, box_w=Inches(4.2), box_h=Inches(0.55), gap=Inches(0.22)):
    y = top
    for i, step in enumerate(steps):
        x = center_x - box_w / 2
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_w, box_h)
        box.fill.solid()
        box.fill.fore_color.rgb = BLUE if i % 2 == 0 else GOLD
        box.line.fill.background()
        box.shadow.inherit = False
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = step
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = WHITE if i % 2 == 0 else DARK
        y += box_h
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, center_x - Inches(0.12), y, Inches(0.24), gap)
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GREY
            arrow.line.fill.background()
            arrow.shadow.inherit = False
            y += gap
    return y


def add_table(slide, left, top, width, height, headers, rows, col_widths=None, font_size=13):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gtable = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table
    if col_widths:
        for i, w in enumerate(col_widths):
            gtable.columns[i].width = w
    for c, h in enumerate(headers):
        cell = gtable.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(font_size)
                r.font.color.rgb = WHITE
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            cell = gtable.cell(r_i, c_i)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r_i % 2 == 1 else ZEBRA
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = DARK
                    run.font.bold = (c_i == 0)
    return gtable


# --------------------------------------------------------------------------- #
# Slides
# --------------------------------------------------------------------------- #
def slide_01_title(prs):
    slide = blank_slide(prs)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE
    band.line.fill.background()
    band.shadow.inherit = False

    monogram = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.87), Inches(0.9), Inches(1.6), Inches(1.6))
    monogram.fill.solid()
    monogram.fill.fore_color.rgb = DARK
    monogram.line.color.rgb = GOLD
    monogram.line.width = Pt(2)
    monogram.shadow.inherit = False
    p = monogram.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "A"
    run.font.size = Pt(48)
    run.font.bold = True
    run.font.color.rgb = GOLD

    add_textbox(slide, Inches(1), Inches(2.8), Inches(11.33), Inches(0.9), "AELON",
                size=50, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1.3), Inches(3.65), Inches(10.73), Inches(1.0),
                "Plateforme intelligente de service apres-vente bancaire basee sur "
                "l'IA generative, le RAG et une architecture Multi-Agents",
                size=17, color=GOLD, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1), Inches(5.9), Inches(11.33), Inches(0.4),
                "Cynthia Sileu Kapnang", size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1), Inches(6.3), Inches(11.33), Inches(0.4),
                "Master Data Engineer & IA Engineer  ·  Capgemini",
                size=13, color=RGBColor(0xC9, 0xD6, 0xEC), align=PP_ALIGN.CENTER)


def slide_02_context(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "CONTEXTE", "Pourquoi AELON ?")
    add_textbox(slide, Inches(0.6), Inches(1.35), Inches(5.8), Inches(0.4),
                "Constats", size=17, bold=True, color=BLUE)
    add_bullets(slide, Inches(0.6), Inches(1.85), Inches(5.8), Inches(3.6), [
        "Explosion des donnees documentaires bancaires",
        "Difficulte a retrouver rapidement l'information",
        "Temps de traitement eleve pour le support",
        "Besoin de reponses fiables et tracables",
    ], size=16)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(2.0), Inches(5.8), Inches(3.2))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LIGHT_LINE
    card.shadow.inherit = False
    add_textbox(slide, Inches(7.15), Inches(2.25), Inches(5.3), Inches(0.4), "Question", size=15, bold=True, color=GOLD)
    add_quote(slide, Inches(7.15), Inches(2.8), Inches(5.3), Inches(2.1),
              "Comment repondre efficacement aux demandes des utilisateurs "
              "tout en valorisant les conversations produites ?", size=18)
    add_footer(slide, n)


def slide_03_data_to_ai(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "FIL CONDUCTEUR", "De la donnee a l'IA")
    steps = ["Donnees brutes", "Data Engineering", "Base de connaissance", "RAG",
              "IA Generative", "Multi-Agents", "Aide a la decision"]
    add_flow_vertical(slide, steps, SLIDE_W / 2, Inches(1.35), box_w=Inches(4.6), box_h=Inches(0.52), gap=Inches(0.14))
    add_quote(slide, Inches(2.3), Inches(6.65), Inches(8.7), Inches(0.55),
              "L'IA n'est que la derniere brique. Tout commence par la donnee.", size=17)
    add_footer(slide, n)


def slide_04_vision(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "VISION", "Vision globale d'AELON")
    add_flow_horizontal(slide, ["Utilisateur", "Orchestrateur", "Agents specialises"],
                         Inches(1.5), left=Inches(2.4), total_width=Inches(8.5), box_h=Inches(0.9), font_size=15)
    down = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, SLIDE_W / 2 - Inches(0.15), Inches(2.55), Inches(0.3), Inches(0.35))
    down.fill.solid()
    down.fill.fore_color.rgb = GREY
    down.line.fill.background()
    down.shadow.inherit = False
    branches = ["Analytics", "Gouvernance", "Evaluation"]
    col_w = Inches(3.6)
    gap = Inches(0.3)
    total = len(branches) * col_w + (len(branches) - 1) * gap
    start = (SLIDE_W - total) / 2
    for i, b in enumerate(branches):
        x = start + i * (col_w + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.1), col_w, Inches(0.75))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BLUE
        box.line.width = Pt(1.5)
        box.shadow.inherit = False
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = b
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = BLUE
    add_quote(slide, Inches(1.5), Inches(4.5), Inches(10.3), Inches(0.6),
              "AELON n'est pas un chatbot mais une plateforme complete.", size=19)
    add_bullets(slide, Inches(1.8), Inches(5.4), Inches(9.7), Inches(1.4), [
        "Chaque conversation nourrit le pilotage analytique, la gouvernance IA et l'evaluation continue.",
    ], size=14)
    add_footer(slide, n)


def slide_05_medallion(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "DATA PLATFORM", "Architecture medaillon")
    steps = ["Sources", "Bronze", "Silver", "Gold", "Embeddings", "Vector Search", "RAG"]
    bottom = add_flow_horizontal(slide, steps, Inches(1.35), box_h=Inches(0.75), font_size=11)
    layers = [
        ("Bronze", BLUE, "Donnees brutes"),
        ("Silver", GREY, "Nettoyage, standardisation"),
        ("Gold", GOLD, "Connaissance metier"),
    ]
    top = Inches(2.55)
    col_w = Inches(3.9)
    gap = Inches(0.25)
    for i, (name, color, desc) in enumerate(layers):
        x = Inches(0.6) + i * (col_w + gap)
        head = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, col_w, Inches(0.55))
        head.fill.solid()
        head.fill.fore_color.rgb = color
        head.line.fill.background()
        head.shadow.inherit = False
        p = head.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = name
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = WHITE if name != "Gold" else DARK
        body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top + Inches(0.55), col_w, Inches(1.6))
        body.fill.solid()
        body.fill.fore_color.rgb = WHITE
        body.line.color.rgb = LIGHT_LINE
        body.shadow.inherit = False
        add_textbox(slide, x + Inches(0.2), top + Inches(0.75), col_w - Inches(0.4), Inches(1.2),
                    desc, size=14, color=GREY)
    add_quote(slide, Inches(0.6), Inches(5.1), Inches(12.1), Inches(0.8),
              "Chaque couche fiabilise la donnee avant qu'elle n'atteigne l'IA.", size=16)
    add_footer(slide, n)


def slide_06_tables(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "DATA PLATFORM", "Tables Databricks")
    headers = ["Table", "Role"]
    rows = [
        ("bronze_documents", "Ingestion des documents bruts et tracabilite d'origine"),
        ("silver_documents", "Nettoyage, normalisation et structuration des contenus"),
        ("gold_documents", "Corpus documentaire consolide pour le RAG"),
        ("gold_embeddings", "Vecteurs semantiques pour la recherche"),
        ("gold_conversations", "Historique des interactions pour analytics, gouvernance et evaluation"),
    ]
    add_table(slide, Inches(0.9), Inches(1.6), Inches(11.5), Inches(3.4), headers, rows,
              col_widths=[Inches(3.2), Inches(8.3)], font_size=14)
    add_footer(slide, n)


def slide_07_stack(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "TECHNOLOGIE", "Socle technique")
    items = ["Databricks", "Delta Lake", "Unity Catalog", "Vector Search",
             "Azure OpenAI", "FastAPI", "LangGraph"]
    cols = 4
    box_w = Inches(2.75)
    box_h = Inches(1.3)
    gap_x = Inches(0.25)
    gap_y = Inches(0.3)
    total_w = cols * box_w + (cols - 1) * gap_x
    start_x = (SLIDE_W - total_w) / 2
    top = Inches(1.8)
    for i, item in enumerate(items):
        row, col = divmod(i, cols)
        x = start_x + col * (box_w + gap_x)
        y = top + row * (box_h + gap_y)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_w, box_h)
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BLUE
        box.line.width = Pt(1.25)
        box.shadow.inherit = False
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = item
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = BLUE
    add_footer(slide, n)


def slide_08_pipeline(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "PIPELINE", "Pipeline documentaire")
    add_flow_horizontal(slide, ["PDF", "Extraction", "Chunking", "Embeddings", "Vector Search"],
                         Inches(1.7), box_h=Inches(0.9), font_size=13)
    add_textbox(slide, Inches(0.6), Inches(3.2), Inches(6), Inches(0.4),
                "Pourquoi chunker ?", size=17, bold=True, color=BLUE)
    add_bullets(slide, Inches(0.6), Inches(3.7), Inches(11.5), Inches(1.6), [
        "Reduire la taille des documents",
        "Ameliorer la recherche",
    ], size=16)
    add_footer(slide, n)


def slide_09_embeddings(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "RECHERCHE SEMANTIQUE", "Embeddings et Vector Search")
    steps = ["Question utilisateur", "Embedding", "Vector Search", "Top 3 Chunks"]
    add_flow_vertical(slide, steps, SLIDE_W / 2, Inches(1.5), box_w=Inches(4.8), box_h=Inches(0.6), gap=Inches(0.25))
    add_quote(slide, Inches(1.6), Inches(5.7), Inches(10.1), Inches(0.8),
              "L'embedding transforme le texte en vecteur afin de comparer le sens et non les mots.", size=17)
    add_footer(slide, n)


def slide_10_rag(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "GENERATION", "Fonctionnement du RAG")
    add_flow_horizontal(slide, ["Question", "Retrieval", "Contexte", "LLM", "Reponse"],
                         Inches(1.7), box_h=Inches(0.9), font_size=13)
    add_bullets(slide, Inches(0.6), Inches(3.3), Inches(11.5), Inches(2.2), [
        "Reduction des hallucinations",
        "Tracabilite",
        "Contexte metier",
    ], size=16)
    add_footer(slide, n)


def slide_11_langgraph(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "ORCHESTRATION", "Pourquoi LangGraph ?")
    add_flow_horizontal(slide, ["Question", "LangGraph", "Agents specialises", "Reponse"],
                         Inches(1.7), box_h=Inches(0.9), font_size=13, total_width=Inches(10.5), left=Inches(1.4))
    add_textbox(slide, Inches(0.6), Inches(3.3), Inches(6), Inches(0.4), "Role", size=17, bold=True, color=BLUE)
    add_bullets(slide, Inches(0.6), Inches(3.8), Inches(11.5), Inches(2.0), [
        "Gestion des flux",
        "Coordination des agents",
        "Etat de conversation",
    ], size=16)
    add_footer(slide, n)


def slide_12_agents(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "MULTI-AGENTS", "Architecture Multi-Agents")
    headers = ["Agent", "Role"]
    rows = [
        ("Privacy", "Protection des donnees"),
        ("Sentiment", "Ton et langue"),
        ("Fraud", "Detection de fraude"),
        ("Retrieval", "Recherche documentaire"),
        ("Compliance", "Conformite"),
        ("Governance", "Gouvernance"),
        ("Analytics", "Statistiques"),
        ("Evaluation", "KPI"),
    ]
    add_table(slide, Inches(2.4), Inches(1.5), Inches(8.5), Inches(5.4), headers, rows,
              col_widths=[Inches(3.0), Inches(5.5)], font_size=14)
    add_footer(slide, n)


def slide_13_journey(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "CAS D'USAGE", "Parcours complet d'une question")
    add_quote(slide, Inches(0.6), Inches(1.35), Inches(12.1), Inches(0.55),
              "« Comment faire opposition a ma carte bancaire ? »", size=18)
    steps = ["Utilisateur", "Orchestrator", "Retrieval Agent", "Vector Search"]
    steps2 = ["Contexte", "LLM", "Reponse", "gold_conversations"]
    add_flow_horizontal(slide, steps, Inches(2.35), box_h=Inches(0.8), font_size=11.5)
    add_flow_horizontal(slide, steps2, Inches(3.5), box_h=Inches(0.8), font_size=11.5)
    down = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, SLIDE_W / 2 - Inches(0.15), Inches(3.18), Inches(0.3), Inches(0.28))
    down.fill.solid()
    down.fill.fore_color.rgb = GREY
    down.line.fill.background()
    down.shadow.inherit = False
    add_bullets(slide, Inches(0.6), Inches(4.7), Inches(12.1), Inches(1.6), [
        "Chaque etape est tracee, du besoin client jusqu'a la persistance analytique.",
    ], size=15)
    add_footer(slide, n)


def _dashboard_slide(prs, n, kicker, title, bullets, placeholder_label):
    slide = blank_slide(prs)
    add_header(slide, kicker, title)
    add_bullets(slide, Inches(0.6), Inches(1.5), Inches(4.6), Inches(3.5), bullets, size=16)
    add_placeholder(slide, Inches(5.5), Inches(1.5), Inches(7.2), Inches(4.9), placeholder_label)
    add_footer(slide, n)
    return slide


def slide_14_analytics(prs, n):
    _dashboard_slide(prs, n, "PILOTAGE", "Analytics Dashboard",
                      ["Volume de conversations", "Categories", "Tendances"],
                      "Capture d'ecran — Analytics Dashboard")


def slide_15_governance(prs, n):
    _dashboard_slide(prs, n, "PILOTAGE", "Governance Dashboard",
                      ["Qualite des reponses", "Utilisation des sources", "Monitoring des agents"],
                      "Capture d'ecran — Governance Dashboard")


def slide_16_evaluation(prs, n):
    _dashboard_slide(prs, n, "PILOTAGE", "Evaluation Dashboard",
                      ["Retrieval Success Rate", "Source Match Rate", "Category Match Rate", "Keyword Match Rate"],
                      "Capture d'ecran — Evaluation Dashboard")


def slide_17_results(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "RESULTATS", "Resultats de l'evaluation")
    kpis = [
        ("100 %", "Retrieval Success Rate"),
        ("100 %", "Source Match Rate"),
        ("100 %", "Category Coverage"),
        ("62,5 %", "Keyword Match Rate"),
        ("50 %", "Category Match Rate"),
    ]
    n_k = len(kpis)
    card_w = Inches(2.2)
    gap = Inches(0.2)
    total_w = n_k * card_w + (n_k - 1) * gap
    start_x = (SLIDE_W - total_w) / 2
    for i, (value, label) in enumerate(kpis):
        x = start_x + i * (card_w + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.45), card_w, Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BLUE
        card.line.width = Pt(1.25)
        card.shadow.inherit = False
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = value
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = BLUE
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = label
        run2.font.size = Pt(11)
        run2.font.color.rgb = GREY

    add_textbox(slide, Inches(0.8), Inches(3.5), Inches(5.6), Inches(0.4), "Forces", size=17, bold=True, color=GREEN)
    add_bullets(slide, Inches(0.8), Inches(4.0), Inches(5.6), Inches(1.6), [
        "Bonne recuperation documentaire",
        "Bonnes sources",
    ], size=15, header_color=GREEN)

    add_textbox(slide, Inches(6.9), Inches(3.5), Inches(5.6), Inches(0.4),
                "Axes d'amelioration", size=17, bold=True, color=AMBER)
    add_bullets(slide, Inches(6.9), Inches(4.0), Inches(5.6), Inches(1.6), [
        "Categorisation",
        "Latence",
    ], size=15, header_color=AMBER)
    add_footer(slide, n)


def slide_18_conclusion(prs, n):
    slide = blank_slide(prs)
    add_header(slide, "CONCLUSION", "Conclusion & Perspectives")
    add_textbox(slide, Inches(0.6), Inches(1.35), Inches(5.8), Inches(0.4),
                "Ce que demontre AELON", size=16, bold=True, color=BLUE)
    add_bullets(slide, Inches(0.6), Inches(1.85), Inches(5.8), Inches(3.0), [
        "IA generative", "RAG", "LangGraph", "Multi-Agents",
        "Databricks", "Gouvernance", "Analytics",
    ], size=15, header_color=GREEN)
    add_textbox(slide, Inches(6.9), Inches(1.35), Inches(5.8), Inches(0.4),
                "Perspectives", size=16, bold=True, color=BLUE)
    add_bullets(slide, Inches(6.9), Inches(1.85), Inches(5.8), Inches(3.0), [
        "Amelioration des performances",
        "Feedback utilisateur",
        "Guardrails",
        "IA plus explicable",
        "Copilots metier",
    ], size=15)
    add_quote(slide, Inches(1.6), Inches(5.6), Inches(10.1), Inches(0.9),
              "Donnee -> Connaissance -> IA -> Decision", size=22)
    add_footer(slide, n)


# --------------------------------------------------------------------------- #
def build() -> Path:
    prs = new_presentation()
    slide_01_title(prs)
    slide_02_context(prs, 2)
    slide_03_data_to_ai(prs, 3)
    slide_04_vision(prs, 4)
    slide_05_medallion(prs, 5)
    slide_06_tables(prs, 6)
    slide_07_stack(prs, 7)
    slide_08_pipeline(prs, 8)
    slide_09_embeddings(prs, 9)
    slide_10_rag(prs, 10)
    slide_11_langgraph(prs, 11)
    slide_12_agents(prs, 12)
    slide_13_journey(prs, 13)
    slide_14_analytics(prs, 14)
    slide_15_governance(prs, 15)
    slide_16_evaluation(prs, 16)
    slide_17_results(prs, 17)
    slide_18_conclusion(prs, 18)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


if __name__ == "__main__":
    out = build()
    print(f"Presentation generated: {out}")
