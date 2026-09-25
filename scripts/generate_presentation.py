# -*- coding: utf-8 -*-
"""Generate the AELON PowerPoint presentation deck from static content."""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# AELON brandbook palette
BLUE = RGBColor(0x0A, 0x3D, 0x91)
GOLD = RGBColor(0xD4, 0xAF, 0x37)
GREY = RGBColor(0x66, 0x70, 0x85)
BG = RGBColor(0xF5, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1B, 0x22, 0x33)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

OUTPUT_PATH = Path(__file__).resolve().parent.parent / "docs" / "AELON_Presentation.pptx"


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_slide(prs: Presentation):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send background to back
    spTree = slide.shapes._spTree
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)
    return slide


def add_textbox(slide, left, top, width, height, text, size=18, color=DARK,
                 bold=False, align=PP_ALIGN.LEFT, font="Segoe UI", anchor=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return box


def add_header(slide, kicker, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    bar.shadow.inherit = False
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(0.42), Inches(0.3), Inches(0.3))
    dot.fill.solid()
    dot.fill.fore_color.rgb = GOLD
    dot.line.fill.background()
    dot.shadow.inherit = False
    add_textbox(slide, Inches(0.95), Inches(0.08), Inches(11.5), Inches(0.35), kicker,
                size=13, color=GOLD, bold=True)
    add_textbox(slide, Inches(0.95), Inches(0.4), Inches(11.5), Inches(0.65), title,
                size=26, color=WHITE, bold=True)


def add_footer(slide, page_no):
    add_textbox(slide, Inches(0.5), Inches(7.12), Inches(6), Inches(0.3),
                "AELON — Plateforme IA de service apres-vente bancaire", size=10, color=GREY)
    add_textbox(slide, Inches(12.3), Inches(7.12), Inches(0.6), Inches(0.3),
                str(page_no), size=10, color=GREY, align=PP_ALIGN.RIGHT)


def add_bullets(slide, left, top, width, height, items, size=16, color=DARK,
                 bullet_color=BLUE, line_spacing=1.25):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        p.line_spacing = line_spacing
        p.space_after = Pt(8)
        run = p.add_run()
        prefix = "• " if level == 0 else "— "
        run.text = f"{prefix}{text}"
        run.font.size = Pt(size if level == 0 else size - 2)
        run.font.color.rgb = color if level == 0 else GREY
        run.font.name = "Segoe UI"
    return box


def add_title_slide(prs: Presentation):
    slide = blank_slide(prs)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE
    band.line.fill.background()
    band.shadow.inherit = False

    monogram = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.87), Inches(1.15), Inches(1.6), Inches(1.6))
    monogram.fill.solid()
    monogram.fill.fore_color.rgb = DARK
    monogram.line.color.rgb = GOLD
    monogram.line.width = Pt(2)
    monogram.shadow.inherit = False
    tf = monogram.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "A"
    run.font.size = Pt(48)
    run.font.bold = True
    run.font.color.rgb = GOLD

    add_textbox(slide, Inches(1), Inches(3.05), Inches(11.33), Inches(1.0), "AELON",
                size=54, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1), Inches(3.85), Inches(11.33), Inches(0.6),
                "Plateforme intelligente de service apres-vente bancaire",
                size=20, color=GOLD, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1), Inches(4.4), Inches(11.33), Inches(0.5),
                "Data Engineering · GenAI · RAG · Architecture Multi-Agents",
                size=15, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1), Inches(6.4), Inches(11.33), Inches(0.5),
                "From Customer Conversations to Intelligent Decisions",
                size=13, color=RGBColor(0xC9, 0xD6, 0xEC), align=PP_ALIGN.CENTER, bold=True)


def add_overview_slide(prs: Presentation, page_no):
    slide = blank_slide(prs)
    add_header(slide, "PRESENTATION", "Vue d'ensemble")
    add_textbox(slide, Inches(0.6), Inches(1.45), Inches(12.1), Inches(1.0),
                "AELON est developpee dans le cadre du projet RAISE. Elle combine Data Engineering, "
                "IA Generative, Retrieval-Augmented Generation (RAG) et Architecture Multi-Agents pour "
                "assister les utilisateurs bancaires et piloter la performance des agents IA.",
                size=15, color=DARK)
    add_bullets(slide, Inches(0.6), Inches(2.6), Inches(12.1), Inches(3.8), [
        "Depasse le cadre d'un simple chatbot : recherche documentaire, generation de reponses, qualite et analytique.",
        "Assistant bancaire conversationnel ancre sur des sources metier fiables (RAG).",
        "Orchestration d'agents specialises : retrieval, raisonnement, controle, supervision.",
        "Fondations Data sur Databricks Lakehouse (ingestion, transformation, historisation).",
        "Recherche semantique par embeddings via Databricks Vector Search.",
        "Deux tableaux de bord de pilotage : Analytics et Governance.",
        "Cadre d'evaluation IA (RAG + agents) sur un dataset metier.",
    ], size=15)
    add_footer(slide, page_no)


def add_features_slide(prs: Presentation, page_no):
    slide = blank_slide(prs)
    add_header(slide, "CAPACITES", "Fonctionnalites cles")
    features = [
        ("Banking Chat", "Assistant bancaire conversationnel alimente par IA Generative"),
        ("Architecture Multi-Agents", "Orchestration d'agents specialises (retrieval, raisonnement, controle)"),
        ("Databricks Lakehouse", "Ingestion, transformation, historisation et exploitation analytique"),
        ("Databricks Vector Search", "Recherche semantique par embeddings"),
        ("RAG", "Reponses enrichies et ancrees sur un corpus documentaire bancaire"),
        ("Analytics Dashboard", "Volumes, categories et tendances conversationnelles"),
        ("Governance Dashboard", "Qualite, usage des sources et monitoring des agents"),
        ("AI Evaluation Framework", "Mesure des performances RAG et agents sur un dataset metier"),
    ]
    top = Inches(1.5)
    col_w = Inches(5.95)
    row_h = Inches(1.35)
    for i, (name, desc) in enumerate(features):
        col = i % 2
        row = i // 2
        left = Inches(0.6) + col * (col_w + Inches(0.3))
        y = top + row * (row_h + Inches(0.12))
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, y, col_w, row_h)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(0xDD, 0xE3, 0xEC)
        card.line.width = Pt(1)
        card.shadow.inherit = False
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, y, Inches(0.08), row_h)
        accent.fill.solid()
        accent.fill.fore_color.rgb = GOLD
        accent.line.fill.background()
        accent.shadow.inherit = False
        add_textbox(slide, left + Inches(0.25), y + Inches(0.1), col_w - Inches(0.4), Inches(0.35),
                    name, size=14, bold=True, color=BLUE)
        add_textbox(slide, left + Inches(0.25), y + Inches(0.5), col_w - Inches(0.4), Inches(0.75),
                    desc, size=11.5, color=GREY)
    add_footer(slide, page_no)


def add_architecture_slide(prs: Presentation, page_no):
    slide = blank_slide(prs)
    add_header(slide, "ARCHITECTURE", "Architecture globale end-to-end")
    steps = [
        "Utilisateur", "Banking Chat", "Orchestrateur", "Architecture Multi-Agents",
        "Retrieval Agent", "Databricks Vector Search", "Corpus documentaire",
        "Modele generatif", "Reponse",
    ]
    left = Inches(0.5)
    top = Inches(1.7)
    box_w = Inches(1.28)
    box_h = Inches(0.85)
    gap = Inches(0.12)
    for i, step in enumerate(steps):
        x = left + i * (box_w + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, box_w, box_h)
        box.fill.solid()
        box.fill.fore_color.rgb = BLUE if i % 2 == 0 else GOLD
        box.line.fill.background()
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = step
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = WHITE if i % 2 == 0 else DARK
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + box_w, top + Inches(0.28),
                                            gap, Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GREY
            arrow.line.fill.background()
            arrow.shadow.inherit = False

    add_textbox(slide, Inches(0.5), Inches(2.9), Inches(12.3), Inches(0.4),
                "La reponse alimente en continu la boucle de pilotage :", size=15, bold=True, color=DARK)

    loop = ["gold_conversations", "Analytics", "Governance", "Evaluation"]
    lw = Inches(2.6)
    lh = Inches(0.8)
    lgap = Inches(0.3)
    total = len(loop) * lw + (len(loop) - 1) * lgap
    lstart = (SLIDE_W - total) / 2
    for i, label in enumerate(loop):
        x = lstart + i * (lw + lgap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.5), lw, lh)
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BLUE
        box.line.width = Pt(1.5)
        box.shadow.inherit = False
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = BLUE
        if i < len(loop) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + lw, Inches(3.75), lgap, Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GREY
            arrow.line.fill.background()
            arrow.shadow.inherit = False

    add_textbox(slide, Inches(0.6), Inches(4.7), Inches(12.1), Inches(1.6),
                "Chaque interaction utilisateur devient a la fois une assistance immediate et une source "
                "de pilotage pour l'amelioration continue du systeme bancaire intelligent.",
                size=14, color=GREY)
    add_footer(slide, page_no)


def add_medallion_slide(prs: Presentation, page_no):
    slide = blank_slide(prs)
    add_header(slide, "DATA PLATFORM", "Architecture medaillon (Lakehouse)")
    layers = [
        ("Bronze", BLUE, ["bronze_documents"],
         "Ingestion des documents bruts, conservation des traces d'origine, reprise fiable des traitements."),
        ("Silver", GREY, ["silver_documents"],
         "Nettoyage, normalisation et structuration des contenus documentaires."),
        ("Gold", GOLD, ["gold_documents", "gold_embeddings", "gold_conversations"],
         "Couche de consommation pour le RAG, la recherche vectorielle, l'analytics et la gouvernance."),
    ]
    top = Inches(1.6)
    col_w = Inches(3.9)
    gap = Inches(0.25)
    for i, (name, color, tables, desc) in enumerate(layers):
        x = Inches(0.6) + i * (col_w + gap)
        header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, col_w, Inches(0.6))
        header.fill.solid()
        header.fill.fore_color.rgb = color
        header.line.fill.background()
        header.shadow.inherit = False
        tf = header.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = name
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = WHITE if name != "Gold" else DARK

        body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top + Inches(0.6), col_w, Inches(3.4))
        body.fill.solid()
        body.fill.fore_color.rgb = WHITE
        body.line.color.rgb = RGBColor(0xDD, 0xE3, 0xEC)
        body.shadow.inherit = False

        add_bullets(slide, x + Inches(0.2), top + Inches(0.8), col_w - Inches(0.4), Inches(1.2),
                    [(t, 0) for t in tables], size=13, bullet_color=color)
        add_textbox(slide, x + Inches(0.2), top + Inches(1.9), col_w - Inches(0.4), Inches(1.9),
                    desc, size=12, color=GREY)
        if i < len(layers) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + col_w, top + Inches(1.9), gap, Inches(0.4))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GREY
            arrow.line.fill.background()
            arrow.shadow.inherit = False
    add_textbox(slide, Inches(0.6), Inches(5.3), Inches(12.1), Inches(1.4),
                "Le pipeline documentaire AELON s'appuie sur une architecture medaillon pour garantir "
                "qualite, tracabilite et exploitation industrielle des donnees.", size=14, color=DARK)
    add_footer(slide, page_no)


def add_spaces_slide(prs: Presentation, page_no):
    slide = blank_slide(prs)
    add_header(slide, "PRODUIT", "Espaces produit par role")
    spaces = [
        ("Customer Space", BLUE, "ROLE: CUSTOMER", [
            "Surface principale : Chat uniquement",
            "Historique de conversation, messages, reponses IA",
            "KPI, dashboards et internals des agents masques",
            "Flux privacy-first avant affichage/stockage",
        ]),
        ("Business Analytics Space", RGBColor(0x6C, 0x3C, 0xB5), "ROLE: BUSINESS_ANALYST", [
            "Source : interactions du chat uniquement",
            "Dashboard, Conversations, Fraude, Sentiment, Escalations, KPI",
            "Analytics Copilot repond a partir des metriques d'interaction",
        ]),
        ("Data Governance Space", RGBColor(0x1E, 0x7A, 0x4C), "ROLE: DATA_STEWARD", [
            "Data Governance, Observabilite, Evaluation Agent",
            "AI Governance Copilot : evaluation, observabilite, gouvernance",
            "Focus : IA responsable, qualite des donnees, conformite, risque",
        ]),
    ]
    top = Inches(1.55)
    col_w = Inches(3.95)
    gap = Inches(0.2)
    for i, (name, color, role, items) in enumerate(spaces):
        x = Inches(0.55) + i * (col_w + gap)
        header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, col_w, Inches(0.95))
        header.fill.solid()
        header.fill.fore_color.rgb = color
        header.line.fill.background()
        header.shadow.inherit = False
        tf = header.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = name
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = WHITE
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = role
        run2.font.size = Pt(10)
        run2.font.color.rgb = RGBColor(0xE9, 0xE9, 0xE9)

        body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top + Inches(0.95), col_w, Inches(3.6))
        body.fill.solid()
        body.fill.fore_color.rgb = WHITE
        body.line.color.rgb = RGBColor(0xDD, 0xE3, 0xEC)
        body.shadow.inherit = False
        add_bullets(slide, x + Inches(0.2), top + Inches(1.15), col_w - Inches(0.4), Inches(3.2),
                    items, size=12, bullet_color=color)
    add_footer(slide, page_no)


def add_flow_slide(prs: Presentation, page_no):
    slide = blank_slide(prs)
    add_header(slide, "FONCTIONNEL", "Flux end-to-end")
    steps = [
        "Saisie utilisateur", "Masquage Privacy Agent", "Fraud Agent", "Sentiment Agent",
        "Decision L0 / L1", "Retrieval Agent (contexte RAG)", "Generation GPT-4o / GPT-4.1",
        "Controles Compliance Agent", "Annotation Explainability", "Affichage reponse client",
        "Ingestion Databricks Lakehouse", "Consommation Analytics & Governance", "Metriques Observabilite & Evaluation",
    ]
    left_col = steps[:7]
    right_col = steps[7:]
    for col_i, col in enumerate([left_col, right_col]):
        x = Inches(0.6) + col_i * Inches(6.3)
        y = Inches(1.55)
        for i, step in enumerate(col):
            num = col_i * 7 + i + 1
            circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Inches(0.4), Inches(0.4))
            circle.fill.solid()
            circle.fill.fore_color.rgb = GOLD if num % 2 == 0 else BLUE
            circle.line.fill.background()
            circle.shadow.inherit = False
            tf = circle.text_frame
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = str(num)
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = WHITE if num % 2 != 0 else DARK
            add_textbox(slide, x + Inches(0.55), y - Inches(0.02), Inches(5.4), Inches(0.45),
                        step, size=13, color=DARK)
            y += Inches(0.63)
    add_footer(slide, page_no)


def add_security_slide(prs: Presentation, page_no):
    slide = blank_slide(prs)
    add_header(slide, "TRUST & COMPLIANCE", "Securite, confidentialite et gouvernance IA")
    add_textbox(slide, Inches(0.6), Inches(1.45), Inches(5.8), Inches(0.4),
                "Regles de securite et confidentialite", size=16, bold=True, color=BLUE)
    add_bullets(slide, Inches(0.6), Inches(1.9), Inches(5.8), Inches(2.6), [
        "Seul le contenu masque est affiche, transmis aux agents et persiste",
        "Categories PII prioritaires : telephone, email, IBAN, compte, carte, identite",
        "Metadonnees de conformite et d'explicabilite capturees pour supervision",
    ], size=13)
    add_textbox(slide, Inches(6.9), Inches(1.45), Inches(5.8), Inches(0.4),
                "Signaux LLMOps & Observabilite", size=16, bold=True, color=BLUE)
    add_bullets(slide, Inches(6.9), Inches(1.9), Inches(5.8), Inches(2.6), [
        "Relevance Score", "Faithfulness Score", "Hallucination Rate",
        "Compliance Score", "Answer Quality", "Response time", "Escalation volume",
    ], size=13)
    add_textbox(slide, Inches(0.6), Inches(4.6), Inches(12.1), Inches(0.4),
                "Sources documentaires de reference", size=16, bold=True, color=BLUE)
    add_bullets(slide, Inches(0.6), Inches(5.05), Inches(12.1), Inches(1.6), [
        "Banque de France", "ACPR", "CNIL", "FBF", "FAQ metier",
    ], size=13)
    add_footer(slide, page_no)


def add_techstack_slide(prs: Presentation, page_no):
    slide = blank_slide(prs)
    add_header(slide, "TECHNOLOGIE", "Stack technique")
    rows = [
        ("Data Engineering", "Databricks, PySpark, Delta Lake"),
        ("IA", "GenAI, RAG, Embeddings"),
        ("Recherche", "Databricks Vector Search"),
        ("Gouvernance", "Unity Catalog"),
        ("Backend", "Python (FastAPI, Uvicorn)"),
        ("Frontend", "HTML, CSS, JavaScript"),
        ("Analytics", "Dashboards & KPI"),
    ]
    top = Inches(1.6)
    row_h = Inches(0.68)
    col1_w = Inches(3.6)
    col2_w = Inches(8.5)
    header_row = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), top, col1_w + col2_w, Inches(0.55))
    header_row.fill.solid()
    header_row.fill.fore_color.rgb = BLUE
    header_row.line.fill.background()
    header_row.shadow.inherit = False
    add_textbox(slide, Inches(0.75), top + Inches(0.08), col1_w - Inches(0.2), Inches(0.4),
                "Domaine", size=13, bold=True, color=WHITE)
    add_textbox(slide, Inches(0.6) + col1_w + Inches(0.15), top + Inches(0.08), col2_w - Inches(0.2), Inches(0.4),
                "Technologies", size=13, bold=True, color=WHITE)
    y = top + Inches(0.55)
    for i, (domain, tech) in enumerate(rows):
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), y, col1_w + col2_w, row_h)
        band.fill.solid()
        band.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(0xEE, 0xF1, 0xF6)
        band.line.fill.background()
        band.shadow.inherit = False
        add_textbox(slide, Inches(0.75), y + Inches(0.15), col1_w - Inches(0.2), Inches(0.4),
                    domain, size=13, bold=True, color=BLUE)
        add_textbox(slide, Inches(0.6) + col1_w + Inches(0.15), y + Inches(0.15), col2_w - Inches(0.2), Inches(0.4),
                    tech, size=13, color=DARK)
        y += row_h
    add_footer(slide, page_no)


def add_closing_slide(prs: Presentation):
    slide = blank_slide(prs)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE
    band.line.fill.background()
    band.shadow.inherit = False
    add_textbox(slide, Inches(1), Inches(2.9), Inches(11.33), Inches(0.8),
                "Merci", size=44, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1), Inches(3.75), Inches(11.33), Inches(0.5),
                "Questions & discussion", size=18, color=GOLD, align=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1), Inches(6.6), Inches(11.33), Inches(0.4),
                "AELON — Projet RAISE", size=12, color=RGBColor(0xC9, 0xD6, 0xEC), align=PP_ALIGN.CENTER)


def build() -> Path:
    prs = new_presentation()
    add_title_slide(prs)
    page = 2
    add_overview_slide(prs, page); page += 1
    add_features_slide(prs, page); page += 1
    add_architecture_slide(prs, page); page += 1
    add_medallion_slide(prs, page); page += 1
    add_spaces_slide(prs, page); page += 1
    add_flow_slide(prs, page); page += 1
    add_security_slide(prs, page); page += 1
    add_techstack_slide(prs, page); page += 1
    add_closing_slide(prs)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


if __name__ == "__main__":
    out = build()
    print(f"Presentation generated: {out}")
